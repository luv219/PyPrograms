import os
import csv
from docx import Document
from tkinter import Tk, filedialog
from openpyxl import Workbook

def ask_for_file_path():
    root = Tk()
    root.withdraw()
    file_path = filedialog.askopenfilename(title="Select Word Document", filetypes=[("Word Documents", "*.docx")])
    return file_path

def convert_docx_to_excel(docx_file, excel_file):
    # Open the .docx file
    doc = Document(docx_file)

    # Create a new Excel workbook
    workbook = Workbook()
    sheet = workbook.active

    # Write data to Excel
    for paragraph in doc.paragraphs:
        data = paragraph.text.split(',')
        sheet.append(data)

    # Save the Excel file
    workbook.save(excel_file)
    print(f"Conversion successful. Data written to {excel_file}")

# Ask the user to select a Word document
input_docx_file = ask_for_file_path()

# Generate output Excel file path
script_directory = os.path.dirname(os.path.abspath(__file__))
file_name = os.path.splitext(os.path.basename(input_docx_file))[0]
output_path_excel = os.path.join(script_directory, f'{file_name}_output.xlsx')

# Convert .docx to .xlsx
convert_docx_to_excel(input_docx_file, output_path_excel)
